"""Загрузка и агрегация данных «Совпадение ИИ с заключениями врачей-экспертов»
(определение инвалидности: заочное vs очное освидетельствование).

Данные из data/disability.xlsx (двухуровневая шапка, данные с 3-й строки).
Держим в памяти (≈345 строк region×class) и агрегируем на лету.
"""
import os
import re
from openpyxl import load_workbook

# Столбцы (0-based) в disability.xlsx
# 0 ID · 1 Регион · 2 Класс болезней · 3 Всего · 4 Дети · 5 Взрослые
# 6 Заочка всего · 7 Заочка совпало группа · 8..14 разбивка групп (I,II,III,реб0,реб1,реб2,реб3)
# 15 Заочка срок · 16 Заочка группа+срок
# 17 Очка всего · 18 Очка совпало группа · 19..25 разбивка групп · 26 Очка срок · 27 Очка группа+срок

DISABILITY = []   # список dict по строкам

# нормализованное имя региона -> kato_region нашей карты (id_reg в geojson)
_REGION_KATO = {
    'акмолинская': 11, 'актюбинская': 15, 'алматинская': 19, 'атырауская': 23,
    'зко': 27, 'жамбылская': 31, 'жетісу': 33, 'жетысу': 33, 'карагандинская': 35,
    'костанайская': 39, 'кызылординская': 43, 'мангистауская': 47, 'павлодарская': 55,
    'ско': 59, 'туркестанская': 61, 'ұлытау': 62, 'улытау': 62, 'вко': 63,
    'гастана': 71, 'галматы': 75, 'гшымкент': 79, 'абай': 10,
}


def _norm(s):
    s = str(s or '').lower().replace('область', '').replace('обл.', '').replace('обл', '')
    return re.sub(r'[^а-яёіұғқңөүһәa-z0-9]', '', s)


def ai_region_kato(name):
    return _REGION_KATO.get(_norm(name))


def _i(v):
    try:
        return int(v)
    except (ValueError, TypeError):
        try:
            return int(float(v))
        except (ValueError, TypeError):
            return 0


def load_disability():
    DISABILITY.clear()
    path = os.path.join(os.path.dirname(__file__), "data", "disability.xlsx")
    if not os.path.exists(path):
        print("load_disability: disability.xlsx не найден — пропуск")
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[1] is None and row[0] is None:
            continue
        rname = str(row[1]).strip() if row[1] else None
        disease = str(row[2]).strip() if row[2] else None
        if not rname or not disease:
            continue
        n = [_i(x) for x in row]
        DISABILITY.append({
            "rid": str(row[0]).strip() if row[0] else None,
            "region": rname,
            "kato": ai_region_kato(rname),
            "disease": disease,
            "total": n[3], "deti": n[4], "vzr": n[5],
            "z_tot": n[6], "z_grp": n[7],
            "z_g": [n[8], n[9], n[10], n[11], n[12], n[13], n[14]],
            "z_srok": n[15], "z_full": n[16],
            "o_tot": n[17], "o_grp": n[18],
            "o_g": [n[19], n[20], n[21], n[22], n[23], n[24], n[25]],
            "o_srok": n[26], "o_full": n[27],
        })
    wb.close()
    print(f"load_disability: загружено {len(DISABILITY)} строк, "
          f"регионов: {len({r['kato'] for r in DISABILITY if r['kato']})}")


def _filtered(region_id=None, disease=None):
    rows = DISABILITY
    if region_id is not None:
        rows = [r for r in rows if r["kato"] == region_id]
    if disease:
        rows = [r for r in rows if r["disease"] == disease]
    return rows


def _pct(a, b):
    return round(a / b * 100, 1) if b else 0.0


def summary(region_id=None, disease=None):
    rows = _filtered(region_id, disease)
    S = lambda k: sum(r[k] for r in rows)
    total = S("total")
    z_tot, o_tot = S("z_tot"), S("o_tot")
    z_grp, o_grp = S("z_grp"), S("o_grp")
    z_full, o_full = S("z_full"), S("o_full")
    z_srok, o_srok = S("z_srok"), S("o_srok")
    grp_all = z_grp + o_grp
    full_all = z_full + o_full
    return {
        "total": total, "deti": S("deti"), "vzr": S("vzr"),
        "z_tot": z_tot, "o_tot": o_tot,
        "grp_all": grp_all, "grp_pct": _pct(grp_all, total),
        "grp_z_pct": _pct(z_grp, z_tot), "grp_o_pct": _pct(o_grp, o_tot),
        "full_all": full_all, "full_pct": _pct(full_all, total),
        "full_z_pct": _pct(z_full, z_tot), "full_o_pct": _pct(o_full, o_tot),
        "form_z_pct": _pct(z_tot, total), "form_o_pct": _pct(o_tot, total),
        # уровень совпадения по форме (вкладка «Уровень совпадения»)
        "z_grp_pct": _pct(z_grp, z_tot), "z_srok_pct": _pct(z_srok, z_tot), "z_full_pct": _pct(z_full, z_tot),
        "o_grp_pct": _pct(o_grp, o_tot), "o_srok_pct": _pct(o_srok, o_tot), "o_full_pct": _pct(o_full, o_tot),
        "z_grp": z_grp, "o_grp": o_grp, "z_full": z_full, "o_full": o_full,
        "z_srok": z_srok, "o_srok": o_srok,
    }


def regions(disease=None):
    """Агрегаты по регионам — для рейтинга и раскраски карты."""
    agg = {}
    for r in _filtered(None, disease):
        if not r["kato"]:
            continue
        a = agg.setdefault(r["kato"], {"id_reg": r["kato"], "name": r["region"],
                                       "total": 0, "z_tot": 0, "o_tot": 0,
                                       "grp": 0, "z_grp": 0, "o_grp": 0})
        a["total"] += r["total"]; a["z_tot"] += r["z_tot"]; a["o_tot"] += r["o_tot"]
        a["z_grp"] += r["z_grp"]; a["o_grp"] += r["o_grp"]
        a["grp"] += r["z_grp"] + r["o_grp"]
    out = []
    for a in agg.values():
        a["match_pct"] = _pct(a["grp"], a["total"])
        a["z_match_pct"] = _pct(a["z_grp"], a["z_tot"])
        a["o_match_pct"] = _pct(a["o_grp"], a["o_tot"])
        a["z_pct"] = _pct(a["z_tot"], a["total"])
        out.append(a)
    out.sort(key=lambda x: -x["total"])
    return out


GROUP_LABELS = ["I группа", "II группа", "III группа", "Ребёнок до 7 лет",
                "Ребёнок, I группа", "Ребёнок, II группа", "Ребёнок, III группа"]


def groups(region_id=None, disease=None):
    rows = _filtered(region_id, disease)
    z = [sum(r["z_g"][i] for r in rows) for i in range(7)]
    o = [sum(r["o_g"][i] for r in rows) for i in range(7)]
    return [{"label": GROUP_LABELS[i], "z": z[i], "o": o[i]} for i in range(7)]


def diseases(region_id=None):
    """Таблица по классам болезней."""
    agg = {}
    for r in _filtered(region_id, None):
        a = agg.setdefault(r["disease"], {"disease": r["disease"], "total": 0,
                                          "z_tot": 0, "o_tot": 0, "z_grp": 0, "o_grp": 0,
                                          "z_full": 0, "o_full": 0})
        a["total"] += r["total"]; a["z_tot"] += r["z_tot"]; a["o_tot"] += r["o_tot"]
        a["z_grp"] += r["z_grp"]; a["o_grp"] += r["o_grp"]
        a["z_full"] += r["z_full"]; a["o_full"] += r["o_full"]
    out = []
    for a in agg.values():
        a["grp_z_pct"] = _pct(a["z_grp"], a["z_tot"]); a["grp_o_pct"] = _pct(a["o_grp"], a["o_tot"])
        a["full_z_pct"] = _pct(a["z_full"], a["z_tot"]); a["full_o_pct"] = _pct(a["o_full"], a["o_tot"])
        out.append(a)
    out.sort(key=lambda x: -x["total"])
    return out


def disease_list():
    seen = []
    for r in DISABILITY:
        if r["disease"] not in seen:
            seen.append(r["disease"])
    return seen
