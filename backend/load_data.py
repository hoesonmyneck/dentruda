"""Загрузка данных RFPM из data/data.xlsx в таблицу payments.

Колонки читаются ПО ИМЕНИ ЗАГОЛОВКА. Вид помощи берём из сгруппированного
столбца RFPM_NAME_ГРУПП (не из RFPM_NAME).
"""
import os
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from database import engine, Base, Payment


# ── парсеры значений ─────────────────────────────────────────────────────────
def parse_date(val):
    if val is None:
        return None
    if isinstance(val, (date, datetime)):
        return val if isinstance(val, date) else val.date()
    s = str(val).strip()
    if not s:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_num(val, cast=float):
    if val is None:
        return None
    if isinstance(val, str):
        val = val.replace('\xa0', '').replace(' ', '').replace(',', '.')
        if not val:
            return None
    try:
        return cast(val)
    except (ValueError, TypeError):
        return None


def parse_gender(val):
    """Пол в данных — текст Мужской/Женский; маппим в 1/2."""
    if val is None:
        return None
    s = str(val).strip().upper()
    if s.startswith('М'):
        return 1
    if s.startswith('Ж'):
        return 2
    return parse_num(val, int)


def _s(val):
    if val is None:
        return None
    s = str(val).strip()
    return s or None


# ── парсинг листа ────────────────────────────────────────────────────────────
def parse_rows(ws) -> list:
    """Распарсить лист в список Payment по именам заголовков."""
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    hmap = {str(h).strip().upper(): i for i, h in enumerate(header) if h is not None}

    def col(row, *names):
        for n in names:
            i = hmap.get(n.upper())
            if i is not None and i < len(row):
                return row[i]
        return None

    rows_data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sicid = parse_num(col(row, 'SICID'), int)
        group = _s(col(row, 'RFPM_NAME_ГРУПП'))
        if sicid is None and group is None:
            continue
        rows_data.append(Payment(
            sicid=sicid,
            rfpm_group=group,
            rfpm_id=_s(col(row, 'RFPM_ID')),
            rfpm_name=_s(col(row, 'RFPM_NAME')),
            d_naz=parse_date(col(row, 'D_NAZ')),
            nsum=parse_num(col(row, 'NSUM')),
            gender_id=parse_gender(col(row, 'Пол')),
            vozrast=parse_num(col(row, 'Возраст', 'VOZRAST'), int),
            tzhs=_s(col(row, 'TZHS')),
            kato_region=parse_num(col(row, 'КАТО области', 'KATO_REG'), int),
            kato_raion=parse_num(col(row, 'КАТО района', 'KATO_RAI'), int),
            kato_regname=_s(col(row, 'KATO_REGNAME')),
            kato_rainame=_s(col(row, 'KATO_RAINAME')),
        ))
    return rows_data


def _data_path():
    return os.path.join(os.path.dirname(__file__), "data", "data.xlsx")


def load_data():
    """Идемпотентно: создать таблицы и залить data.xlsx, если payments пуста."""
    Base.metadata.create_all(bind=engine)

    with engine.connect() as conn:
        total = conn.execute(text("SELECT COUNT(*) FROM payments")).scalar() or 0
    if total > 0:
        print(f"load_data: в БД уже {total} строк — пропуск")
        return

    path = _data_path()
    if not os.path.exists(path):
        print("load_data: data/data.xlsx не найден — пропуск")
        return

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_data = parse_rows(ws)
    wb.close()

    if not rows_data:
        print("load_data: в data.xlsx не найдено строк")
        return

    with Session(engine) as session:
        session.add_all(rows_data)
        session.commit()
    print(f"Loaded {len(rows_data)} rows from data.xlsx")


def replace_payments_from_file(file_obj) -> int:
    """Полная перезаливка данных из переданного xlsx (для обновления через интерфейс)."""
    Base.metadata.create_all(bind=engine)
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_data = parse_rows(ws)
    wb.close()
    if not rows_data:
        raise ValueError("В файле не найдено строк данных")
    with Session(engine) as session:
        session.execute(text("DELETE FROM payments"))
        session.add_all(rows_data)
        session.commit()
    return len(rows_data)
